from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from PIL import Image, ImageDraw, ImageFont
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PINK = RGBColor(233, 30, 143)
NAVY = RGBColor(17, 22, 45)
PURPLE = RGBColor(108, 92, 231)
ORANGE = RGBColor(242, 169, 0)
GRAY = RGBColor(246, 247, 251)
MUTED = RGBColor(116, 120, 138)
WHITE = RGBColor(255, 255, 255)

PLATFORM_PREFIX = {"Instagram": "ig", "Facebook": "fb", "TikTok": "tt"}
METRIC_LABELS = {
    "views": "Views",
    "reach": "Alcance",
    "impressions": "Impresiones",
    "interactions": "Interacciones",
    "er": "E.R.",
    "saves": "Guardados",
    "clicks": "Clicks",
}


def blank_slide_layout(prs):
    try:
        return prs.slide_layouts[6]
    except Exception:
        return prs.slide_layouts[-1]


def _norm_bool(v: Any) -> bool:
    return str(v).strip().lower() in {"true", "1", "si", "sí", "yes", "x"}


def _to_float(v: Any, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return default
    mult = 1.0
    if s.upper().endswith("M"):
        mult = 1_000_000.0
    if s.upper().endswith("K"):
        mult = 1_000.0
    clean = re.sub(r"[^0-9,\.\-]", "", s)
    if "," in clean and "." in clean:
        clean = clean.replace(".", "").replace(",", ".")
    elif "," in clean:
        clean = clean.replace(",", ".")
    elif clean.count(".") > 1:
        clean = clean.replace(".", "")
    try:
        return float(clean) * mult
    except Exception:
        return default


def fmt_int(v: Any) -> str:
    try:
        return f"{int(round(float(v))):,}".replace(",", ".")
    except Exception:
        return "—"


def fmt_compact(v: Any) -> str:
    n = _to_float(v, 0)
    if abs(n) >= 1_000_000:
        return f"{n/1_000_000:.1f}M".replace(".", ",")
    if abs(n) >= 1_000:
        return f"{n/1_000:.0f}K".replace(".", ",")
    return fmt_int(n)


def fmt_pct(v: Any) -> str:
    n = _to_float(v, 0)
    if n <= 1:
        n = n * 100
    return f"{n:.1f}%".replace(".", ",")


def normalize_month_label(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def extract_drive_id(url: str) -> Optional[str]:
    if not url:
        return None
    patterns = [r"/file/d/([A-Za-z0-9_-]+)", r"[?&]id=([A-Za-z0-9_-]+)", r"/d/([A-Za-z0-9_-]+)"]
    for pat in patterns:
        m = re.search(pat, url)
        if m:
            return m.group(1)
    return None


def placeholder_image(path: Path, title: str, subtitle: str = "Asset pendiente", size: Tuple[int, int] = (900, 1100), color=(233, 30, 143)) -> Path:
    img = Image.new("RGB", size, (248, 248, 252))
    draw = ImageDraw.Draw(img)
    # soft gradient-ish bands
    draw.rectangle([0, 0, size[0], int(size[1]*0.18)], fill=color)
    draw.rounded_rectangle([55, int(size[1]*0.25), size[0]-55, int(size[1]*0.72)], radius=36, fill=(255,255,255), outline=(230,230,240), width=3)
    try:
        font_title = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 52)
        font_sub = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 32)
    except Exception:
        font_title = font_sub = None
    wrapped = []
    words = str(title or "Imagen").split()
    line = ""
    for w in words:
        test = (line + " " + w).strip()
        if len(test) > 22:
            wrapped.append(line)
            line = w
        else:
            line = test
    if line:
        wrapped.append(line)
    y = int(size[1]*0.36)
    for line in wrapped[:3]:
        draw.text((95, y), line, fill=(17,22,45), font=font_title)
        y += 62
    draw.text((95, y+35), subtitle, fill=(116,120,138), font=font_sub)
    img.save(path)
    return path


def crop_to_ratio(src: Path, dst: Path, ratio: float) -> Path:
    img = Image.open(src).convert("RGB")
    w, h = img.size
    current = w / h
    if current > ratio:
        new_w = int(h * ratio)
        left = max(0, (w - new_w) // 2)
        img = img.crop((left, 0, left + new_w, h))
    elif current < ratio:
        new_h = int(w / ratio)
        top = max(0, (h - new_h) // 2)
        img = img.crop((0, top, w, top + new_h))
    img.save(dst, quality=92)
    return dst


def download_asset(url: str, out_dir: Path, fallback_title: str, service_account_info: Optional[Dict[str, Any]] = None) -> Path:
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", (fallback_title or "asset")[:50])
    raw = out_dir / f"raw_{safe}.jpg"
    if not url or "..." in str(url):
        return placeholder_image(raw, fallback_title)
    s = str(url).strip()
    # local packaged path support
    local = Path(s)
    if local.exists():
        return local
    try:
        import requests
        headers = {}
        drive_id = extract_drive_id(s)
        if drive_id:
            # Private Drive file via service account when credentials are available.
            if service_account_info:
                from google.oauth2.service_account import Credentials
                from google.auth.transport.requests import Request
                scopes = ["https://www.googleapis.com/auth/drive.readonly"]
                creds = Credentials.from_service_account_info(service_account_info, scopes=scopes)
                creds.refresh(Request())
                headers["Authorization"] = f"Bearer {creds.token}"
                dl_url = f"https://www.googleapis.com/drive/v3/files/{drive_id}?alt=media"
            else:
                dl_url = f"https://drive.google.com/uc?export=download&id={drive_id}"
        else:
            dl_url = s
        r = requests.get(dl_url, headers=headers, timeout=18)
        r.raise_for_status()
        raw.write_bytes(r.content)
        # verify image
        Image.open(raw).verify()
        return raw
    except Exception:
        return placeholder_image(raw, fallback_title, subtitle="No se pudo cargar imagen")


def add_textbox(slide, x, y, w, h, text, size=14, bold=False, color=NAVY, align=None):
    tx = slide.shapes.add_textbox(x, y, w, h)
    tf = tx.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = str(text or "")
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tx


def add_card(slide, x, y, w, h, fill=WHITE, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(242, 243, 248)
    shape.line.width = Pt(0.8)
    return shape


def add_image_cropped(slide, image_path: Path, x, y, w, h, tmpdir: Path):
    ratio = float(w) / float(h)
    dst = tmpdir / f"crop_{len(list(tmpdir.glob('crop_*')))}.jpg"
    crop_to_ratio(image_path, dst, ratio)
    return slide.shapes.add_picture(str(dst), x, y, w, h)


def as_table(ws):
    rows = []
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and v != "" for v in row):
            continue
        rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]})
    return rows


def get_control_value(wb, key: str, default=None):
    if "00_Control" not in wb.sheetnames:
        return default
    for r in as_table(wb["00_Control"]):
        if str(r.get("Campo") or "").strip() == key:
            return r.get("Valor") if r.get("Valor") not in [None, ""] else default
    return default


def build_top3_media(wb, month: str, warnings: List[str]) -> Dict[str, List[Dict[str, Any]]]:
    out = {"Instagram": [], "Facebook": [], "TikTok": []}
    if "01_Content_Raw" not in wb.sheetnames:
        warnings.append("Falta 01_Content_Raw para Top 3 con imágenes.")
        return out
    rows = as_table(wb["01_Content_Raw"])
    # optional note override
    notes = {}
    if "23_Content_Notes" in wb.sheetnames:
        for r in as_table(wb["23_Content_Notes"]):
            if _norm_bool(r.get("use_in_top3")):
                notes[str(r.get("content_id") or "")] = r.get("short_note")
    for platform in out.keys():
        prefix = PLATFORM_PREFIX[platform]
        metric = str(get_control_value(wb, f"{prefix}_top3_metric", "views") or "views").strip().lower()
        secondary = str(get_control_value(wb, f"{prefix}_top3_secondary_metric", "interactions") or "interactions").strip().lower()
        requires_asset = _norm_bool(get_control_value(wb, "top3_requires_asset", True))
        candidates = []
        for r in rows:
            if normalize_month_label(r.get("month")) != normalize_month_label(month):
                continue
            if str(r.get("platform") or "").strip().lower() != platform.lower():
                continue
            if not _norm_bool(r.get("include_top3")):
                continue
            if requires_asset and not r.get("asset_url"):
                continue
            r = dict(r)
            r["rank_metric"] = metric
            r["rank_metric_label"] = METRIC_LABELS.get(metric, metric)
            r["rank_value"] = _to_float(r.get(metric), 0)
            r["secondary_value"] = _to_float(r.get(secondary), 0)
            r["short_note"] = notes.get(str(r.get("content_id") or ""), r.get("short_note") or "")
            candidates.append(r)
        candidates.sort(key=lambda r: (_to_float(r.get(metric), 0), _to_float(r.get(secondary), 0)), reverse=True)
        out[platform] = candidates[:3]
        if len(out[platform]) < 3:
            warnings.append(f"Top 3 {platform}: solo {len(out[platform])}/3 piezas activas con asset_url.")
    return out


def build_asset_rows(wb, sheet: str, month: str, include_month=True) -> List[Dict[str, Any]]:
    if sheet not in wb.sheetnames:
        return []
    rows = []
    for r in as_table(wb[sheet]):
        if include_month and normalize_month_label(r.get("month")) != normalize_month_label(month):
            continue
        if not _norm_bool(r.get("active", True)):
            continue
        rows.append(r)
    return rows


def enrich_media_context(wb, ctx: Dict[str, Any], month: str, ctrl: Dict[str, Any], warnings: List[str]) -> Dict[str, Any]:
    enable_media = _norm_bool(ctrl.get("enable_media_slides") or ctrl.get("use_assets_ppt") or False)
    ctx["_ENABLE_MEDIA_SLIDES"] = enable_media
    if not enable_media:
        ctx["_TOP3_MEDIA"] = {}
        ctx["_SQUAD_ASSETS"] = []
        ctx["_MMPP_ASSETS"] = []
        ctx["_COMPETITION_ASSETS"] = []
        return ctx
    ctx["_TOP3_MEDIA"] = build_top3_media(wb, month, warnings)
    ctx["_SQUAD_ASSETS"] = build_asset_rows(wb, "20_Squad_Assets", month, include_month=False)
    ctx["_MMPP_ASSETS"] = build_asset_rows(wb, "21_MMPP_Assets", month, include_month=True)[:2]
    ctx["_COMPETITION_ASSETS"] = build_asset_rows(wb, "22_Competition_Assets", month, include_month=True)
    return ctx


def slide_bg(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(250, 250, 253)


def add_header(slide, month_upper: str, title: str):
    add_textbox(slide, Inches(0.35), Inches(0.18), Inches(1.6), Inches(0.25), month_upper, size=10, bold=True, color=PINK)
    add_textbox(slide, Inches(0.35), Inches(0.55), Inches(9.2), Inches(0.55), title, size=26, bold=True, color=NAVY)
    add_textbox(slide, Inches(12.4), Inches(0.25), Inches(0.9), Inches(0.25), "maicao+", size=16, bold=True, color=PINK, align=PP_ALIGN.RIGHT)


def append_top3_slide(prs, context, platform: str, tmpdir: Path, service_account_info: Optional[Dict[str, Any]] = None):
    slide = prs.slides.add_slide(blank_slide_layout(prs))
    slide_bg(slide)
    month = context.get("MES_UPPER", "MES")
    add_header(slide, month, f"Top 3 {platform}: contenidos destacados")
    data = (context.get("_TOP3_MEDIA") or {}).get(platform, [])
    colors = [PINK, NAVY, PURPLE]
    card_w, card_h = Inches(3.75), Inches(5.35)
    start_x = Inches(0.45)
    gap = Inches(0.42)
    for i in range(3):
        x = start_x + i * (card_w + gap)
        y = Inches(1.45)
        add_card(slide, x, y, card_w, card_h, WHITE)
        if i < len(data):
            item = data[i]
            title = item.get("content_title") or f"Pieza {i+1}"
            add_textbox(slide, x + Inches(0.22), y + Inches(0.15), card_w - Inches(0.44), Inches(0.45), title, size=14, bold=True, color=NAVY)
            img = download_asset(str(item.get("asset_url") or ""), tmpdir, title, service_account_info)
            add_image_cropped(slide, img, x + Inches(0.45), y + Inches(0.72), card_w - Inches(0.9), Inches(2.25), tmpdir)
            # Note
            note = item.get("short_note") or "Mini lectura pendiente en 23_Content_Notes."
            add_textbox(slide, x + Inches(0.3), y + Inches(3.08), card_w - Inches(0.6), Inches(0.88), note, size=8.8, color=NAVY)
            # Metrics
            mx = x + Inches(0.25)
            my = y + Inches(4.15)
            add_textbox(slide, mx, my, Inches(0.95), Inches(0.22), "Views", size=7.8, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
            add_textbox(slide, mx, my + Inches(0.25), Inches(0.95), Inches(0.3), fmt_compact(item.get("views")), size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
            add_textbox(slide, mx + Inches(1.2), my, Inches(0.95), Inches(0.22), "Alcance", size=7.8, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
            add_textbox(slide, mx + Inches(1.2), my + Inches(0.25), Inches(0.95), Inches(0.3), fmt_compact(item.get("reach")), size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
            add_textbox(slide, mx + Inches(2.4), my, Inches(0.95), Inches(0.22), "Inter.", size=7.8, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
            add_textbox(slide, mx + Inches(2.4), my + Inches(0.25), Inches(0.95), Inches(0.3), fmt_compact(item.get("interactions")), size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
            add_textbox(slide, x + Inches(0.25), y + Inches(4.88), card_w - Inches(0.5), Inches(0.3), f"E.R. {fmt_pct(item.get('er'))} · ranking por {item.get('rank_metric_label', 'métrica')}", size=11, bold=True, color=colors[i], align=PP_ALIGN.CENTER)
        else:
            add_textbox(slide, x + Inches(0.3), y + Inches(2.2), card_w - Inches(0.6), Inches(0.6), "Pieza pendiente", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    return slide


def append_mmpp_assets_slide(prs, context, tmpdir: Path, service_account_info: Optional[Dict[str, Any]] = None):
    rows = context.get("_MMPP_ASSETS") or []
    if not rows:
        return None
    slide = prs.slides.add_slide(blank_slide_layout(prs))
    slide_bg(slide)
    add_header(slide, context.get("MES_UPPER", "MES"), "MMPP: referencias visuales del mes")
    add_textbox(slide, Inches(0.55), Inches(1.35), Inches(5.6), Inches(0.5), "Piezas destacadas para reforzar calidad, precio y valor práctico", size=16, bold=True, color=NAVY)
    for i in range(2):
        x = Inches(1.25 + i*5.8)
        y = Inches(1.95)
        add_card(slide, x, y, Inches(4.7), Inches(4.8))
        if i < len(rows):
            r = rows[i]
            title = r.get("title") or f"Referencia {i+1}"
            img = download_asset(str(r.get("image_url") or ""), tmpdir, title, service_account_info)
            add_image_cropped(slide, img, x + Inches(0.5), y + Inches(0.4), Inches(3.7), Inches(3.1), tmpdir)
            add_textbox(slide, x + Inches(0.45), y + Inches(3.7), Inches(3.8), Inches(0.35), title, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
            add_textbox(slide, x + Inches(0.45), y + Inches(4.12), Inches(3.8), Inches(0.45), r.get("note") or "Referencia visual del mes.", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    return slide


def append_competition_assets_slide(prs, context, tmpdir: Path, service_account_info: Optional[Dict[str, Any]] = None):
    rows = context.get("_COMPETITION_ASSETS") or []
    if not rows:
        return None
    slide = prs.slides.add_slide(blank_slide_layout(prs))
    slide_bg(slide)
    add_header(slide, context.get("MES_UPPER", "MES"), "Competencia: acciones visuales destacadas")
    # group brands
    brands = []
    for r in rows:
        brand = r.get("brand") or "Marca"
        if brand not in brands:
            brands.append(brand)
    brands = brands[:4]
    card_w = Inches(3.0)
    for bi, brand in enumerate(brands):
        x = Inches(0.45 + bi * 3.2)
        y = Inches(1.45)
        add_card(slide, x, y, card_w, Inches(5.25))
        add_textbox(slide, x + Inches(0.22), y + Inches(0.16), card_w - Inches(0.44), Inches(0.32), brand, size=15, bold=True, color=PINK if bi==0 else NAVY, align=PP_ALIGN.CENTER)
        br = [r for r in rows if (r.get("brand") or "") == brand][:3]
        for i, r in enumerate(br):
            img = download_asset(str(r.get("image_url") or ""), tmpdir, f"{brand} {i+1}", service_account_info)
            img_y = y + Inches(0.68 + i*1.42)
            add_image_cropped(slide, img, x + Inches(0.28), img_y, Inches(1.05), Inches(1.15), tmpdir)
            add_textbox(slide, x + Inches(1.45), img_y + Inches(0.05), Inches(1.25), Inches(0.28), r.get("display_group") or "Acción", size=8.5, bold=True, color=MUTED)
            add_textbox(slide, x + Inches(1.45), img_y + Inches(0.34), Inches(1.25), Inches(0.55), r.get("title") or "Pieza", size=9.2, bold=True, color=NAVY)
    return slide


def append_media_slides(prs, context: Dict[str, Any], service_account_info: Optional[Dict[str, Any]] = None) -> int:
    if not context.get("_ENABLE_MEDIA_SLIDES"):
        return 0
    with tempfile.TemporaryDirectory() as t:
        tmpdir = Path(t)
        before = len(prs.slides)
        append_top3_slide(prs, context, "Instagram", tmpdir, service_account_info)
        append_top3_slide(prs, context, "Facebook", tmpdir, service_account_info)
        append_top3_slide(prs, context, "TikTok", tmpdir, service_account_info)
        append_mmpp_assets_slide(prs, context, tmpdir, service_account_info)
        append_competition_assets_slide(prs, context, tmpdir, service_account_info)
        return len(prs.slides) - before

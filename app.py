"""
Maicao Reporting Studio v12

Professional Streamlit UI for generating the Maicao monthly PPT report from
Google Sheets or an uploaded Excel model, with preview/QA before download.
"""
from __future__ import annotations

import html
import io
import json
import re
import tempfile
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook

from generate_report_from_template import build_context, update_ppt, write_validation_report

ROOT = Path(__file__).resolve().parent
DEFAULT_EXCEL = ROOT / "Maicao_Reporte_Input_Model_v12_MEDIA_BLUEPRINT.xlsx"
DEFAULT_TEMPLATE = ROOT / "template" / "Maicao_Template_Visual_v02.pptx"
DEFAULT_OUTPUT_NAME = "Maicao_Reporte_Mensual_Maicao_v12.pptx"

REQUIRED_SHEETS = [
    "00_Control",
    "01_Content_Raw",
    "02_Audience",
    "03_Paid_Media",
    "04_Squad",
    "05_MMPP",
    "06_Competencia",
    "09_Validaciones",
    "13_Platform_KPIs",
    "15_Qualitative_Texts",
    "16_Action_Plan",
    "20_Squad_Assets",
    "21_MMPP_Assets",
    "22_Competition_Assets",
    "23_Content_Notes",
]

KPI_ORDER = [
    ("Visualizaciones", "OVERVIEW_VIEWS", "OVERVIEW_VIEWS_VAR"),
    ("Alcance", "OVERVIEW_REACH", "OVERVIEW_REACH_VAR"),
    ("Engagement rate", "OVERVIEW_ER", "OVERVIEW_ER_VAR"),
    ("Interacciones", "OVERVIEW_INTERACTIONS", "OVERVIEW_INTERACTIONS_VAR"),
    ("Inversión", "OVERVIEW_INVESTMENT", "OVERVIEW_INVESTMENT_VAR"),
]

PLATFORM_ROWS = [
    ("Instagram", "IG"),
    ("Facebook", "FB"),
    ("TikTok", "TT"),
]

TEXT_PREVIEW_KEYS = [
    ("Dashboard", "Lectura de negocio", "EXEC_BUSINESS_READING"),
    ("Qué cambió", "Decisión recomendada", "CHANGE_DECISION"),
    ("Instagram", "Qué funcionó", "IG_WHAT_WORKED"),
    ("Instagram", "Optimización", "IG_OPTIMIZATION"),
    ("Facebook", "Lectura visual", "FB_VISUAL_READING"),
    ("TikTok", "Principio de contenido", "TT_CONTENT_PRINCIPLE"),
    ("TikTok", "Oportunidad", "TT_OPPORTUNITY"),
    ("Squad", "Aprendizaje", "SQUAD_LEARNING"),
    ("MMPP", "Lectura cualitativa", "MMPP_QUAL_READING"),
    ("Competencia", "Benchmark", "COMP_BENCHMARK"),
]


def escape(v: Any) -> str:
    return html.escape("" if v is None else str(v))


def inject_css() -> None:
    st.markdown(
        """
        <style>
        :root {
            --maicao-pink: #E91E8F;
            --maicao-purple: #6C5CE7;
            --maicao-navy: #11162D;
            --maicao-muted: #74788A;
            --maicao-bg: #F7F7FB;
            --maicao-card: rgba(255,255,255,.92);
            --maicao-border: rgba(17,22,45,.08);
            --maicao-green: #15A36D;
            --maicao-warn: #F2A900;
            --maicao-red: #E5484D;
        }
        .stApp {
            background:
                radial-gradient(circle at 10% 10%, rgba(233,30,143,.10), transparent 24rem),
                radial-gradient(circle at 90% 5%, rgba(108,92,231,.12), transparent 26rem),
                linear-gradient(180deg, #FFFFFF 0%, #F8F8FC 45%, #F4F5FA 100%);
        }
        .block-container { padding-top: 2.2rem; padding-bottom: 4rem; max-width: 1200px; }
        [data-testid="stHeader"] { background: rgba(255,255,255,.70); backdrop-filter: blur(10px); }
        div[data-testid="stToolbar"] { opacity: .65; }
        .hero {
            position: relative;
            overflow: hidden;
            padding: 32px 34px;
            border-radius: 28px;
            background: linear-gradient(135deg, #141A34 0%, #272B6F 56%, #E91E8F 140%);
            color: #fff;
            box-shadow: 0 22px 60px rgba(17,22,45,.20);
            margin-bottom: 1.3rem;
        }
        .hero:after {
            content: "";
            position: absolute;
            width: 280px; height: 280px; border-radius: 50%;
            right: -90px; top: -90px;
            background: rgba(255,255,255,.14);
        }
        .hero .eyebrow {
            font-size: .78rem; letter-spacing: .12em; text-transform: uppercase;
            color: rgba(255,255,255,.72); font-weight: 800; margin-bottom: 12px;
        }
        .hero h1 { font-size: 3.2rem; line-height: 1.02; margin: 0 0 12px 0; font-weight: 900; }
        .hero p { color: rgba(255,255,255,.84); font-size: 1.08rem; max-width: 720px; margin: 0; }
        .pill-row { display:flex; gap:10px; flex-wrap:wrap; margin-top: 20px; }
        .pill {
            display: inline-flex; align-items:center; gap:8px;
            border: 1px solid rgba(255,255,255,.22); background: rgba(255,255,255,.10);
            padding: 8px 12px; border-radius: 999px; color: rgba(255,255,255,.92);
            font-weight: 700; font-size: .88rem;
        }
        .section-title { font-size: 1.35rem; font-weight: 900; color: var(--maicao-navy); margin: 1rem 0 .5rem; }
        .soft-card {
            background: var(--maicao-card);
            border: 1px solid var(--maicao-border);
            border-radius: 22px;
            padding: 20px;
            box-shadow: 0 12px 35px rgba(17,22,45,.06);
            height: 100%;
        }
        .status-card {
            background: rgba(255,255,255,.95);
            border: 1px solid var(--maicao-border);
            border-radius: 18px;
            padding: 16px 16px 14px;
            box-shadow: 0 10px 28px rgba(17,22,45,.05);
        }
        .status-label { color: var(--maicao-muted); text-transform: uppercase; font-size: .72rem; letter-spacing: .08em; font-weight: 900; }
        .status-value { color: var(--maicao-navy); font-size: 1.45rem; font-weight: 900; margin-top: 4px; }
        .status-sub { color: var(--maicao-muted); font-size: .85rem; font-weight: 700; margin-top: 2px; }
        .kpi-card {
            background: #fff;
            border: 1px solid var(--maicao-border);
            border-radius: 18px;
            padding: 16px;
            min-height: 118px;
            box-shadow: 0 10px 24px rgba(17,22,45,.045);
        }
        .kpi-label { color: var(--maicao-muted); text-transform: uppercase; font-size: .72rem; letter-spacing: .08em; font-weight: 900; }
        .kpi-value { color: var(--maicao-pink); font-size: 1.7rem; font-weight: 900; margin-top: 10px; }
        .kpi-sub { color: var(--maicao-muted); font-size: .84rem; margin-top: 2px; font-weight: 700; }
        .mini-note { color: var(--maicao-muted); font-size: .92rem; }
        .success-badge, .warn-badge, .neutral-badge {
            display:inline-flex; align-items:center; gap:7px; padding: 7px 11px; border-radius:999px; font-weight:800; font-size:.85rem;
        }
        .success-badge { background: rgba(21,163,109,.12); color: var(--maicao-green); }
        .warn-badge { background: rgba(242,169,0,.14); color: #A66F00; }
        .neutral-badge { background: rgba(108,92,231,.12); color: var(--maicao-purple); }
        .text-card {
            background: #fff;
            border: 1px solid var(--maicao-border);
            border-radius: 18px;
            padding: 16px;
            min-height: 190px;
            box-shadow: 0 8px 22px rgba(17,22,45,.04);
        }
        .text-section { color: var(--maicao-pink); font-size:.78rem; font-weight:900; text-transform:uppercase; letter-spacing:.08em; }
        .text-title { color: var(--maicao-navy); font-weight:900; font-size:1rem; margin-top:6px; }
        .text-body { color:#20243A; margin-top:10px; font-size:.92rem; line-height:1.42; white-space:pre-line; }
        .media-card { background:#fff; border:1px solid var(--maicao-border); border-radius:18px; padding:14px; box-shadow:0 8px 22px rgba(17,22,45,.04); min-height:280px; }
        .media-title { font-weight:900; color:var(--maicao-navy); font-size:.95rem; margin-top:8px; }
        .media-meta { color:var(--maicao-muted); font-size:.82rem; font-weight:700; margin-top:4px; }
        .help-box { background:#fff; border:1px solid var(--maicao-border); border-left:4px solid var(--maicao-pink); padding:16px 18px; border-radius:16px; }
        .footer-note { text-align:center; color:var(--maicao-muted); margin-top:2rem; font-size:.86rem; }
        div.stButton > button:first-child {
            background: linear-gradient(135deg, var(--maicao-pink), #ff4fa8);
            color: white;
            border: 0;
            border-radius: 14px;
            font-weight: 900;
            padding: .75rem 1.1rem;
            box-shadow: 0 10px 22px rgba(233,30,143,.24);
        }
        div.stDownloadButton > button:first-child {
            background: linear-gradient(135deg, var(--maicao-navy), var(--maicao-purple));
            color: white;
            border: 0;
            border-radius: 14px;
            font-weight: 900;
            padding: .75rem 1.1rem;
            box-shadow: 0 10px 22px rgba(17,22,45,.22);
        }
        div[data-testid="stRadio"] label { font-weight:700; }
        .stTabs [data-baseweb="tab-list"] { gap: 10px; }
        .stTabs [data-baseweb="tab"] {
            background: rgba(255,255,255,.85); border: 1px solid var(--maicao-border);
            border-radius: 999px; padding: 10px 16px; font-weight: 800;
        }
        .stTabs [aria-selected="true"] {
            background: var(--maicao-navy) !important; color: white !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def extract_sheet_id(url_or_id: str) -> str:
    s = (url_or_id or "").strip()
    if not s:
        raise ValueError("Pega el link o ID del Google Sheet.")
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", s)
    if m:
        return m.group(1)
    if re.fullmatch(r"[a-zA-Z0-9-_]{20,}", s):
        return s
    raise ValueError("No pude detectar el ID del Google Sheet. Pega la URL completa o el ID.")


def has_streamlit_secret() -> bool:
    try:
        return "gcp_service_account" in st.secrets
    except Exception:
        return False


def get_service_account_info(uploaded_file: Optional[Any]) -> Dict[str, Any]:
    try:
        if "gcp_service_account" in st.secrets:
            return dict(st.secrets["gcp_service_account"])
    except Exception:
        pass
    if uploaded_file is None:
        raise ValueError("La conexión Google no está configurada. Sube el JSON en configuración avanzada o agrega secrets en Streamlit Cloud.")
    try:
        return json.loads(uploaded_file.getvalue().decode("utf-8"))
    except Exception as exc:
        raise ValueError(f"No pude leer la credencial Google: {exc}")


def google_sheet_to_xlsx(spreadsheet_url_or_id: str, service_account_info: Dict[str, Any], output_path: Path) -> Path:
    try:
        import gspread
        from google.oauth2.service_account import Credentials
    except ImportError as exc:
        raise RuntimeError("Faltan dependencias Google. Instala requirements.txt.") from exc

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets.readonly",
        "https://www.googleapis.com/auth/drive.readonly",
    ]
    creds = Credentials.from_service_account_info(service_account_info, scopes=scopes)
    client = gspread.authorize(creds)
    spreadsheet_id = extract_sheet_id(spreadsheet_url_or_id)
    sh = client.open_by_key(spreadsheet_id)

    wb = Workbook()
    wb.remove(wb.active)
    for worksheet in sh.worksheets():
        title = worksheet.title[:31]
        ws = wb.create_sheet(title=title)
        try:
            values = worksheet.get_all_values(value_render_option="UNFORMATTED_VALUE")
        except TypeError:
            values = worksheet.get_all_values()
        for row_idx, row in enumerate(values, start=1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(output_path)
    return output_path


def save_uploaded_excel(uploaded_file: Any, output_path: Path) -> Path:
    output_path.write_bytes(uploaded_file.getvalue())
    return output_path


def sheet_name_check(xlsx_path: Path) -> list[str]:
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        existing = set(wb.sheetnames)
        return [s for s in REQUIRED_SHEETS if s not in existing]
    except Exception:
        return []


def analyze_xlsx(xlsx_bytes: bytes) -> Tuple[Dict[str, Any], list[str]]:
    with tempfile.TemporaryDirectory() as tmpdir:
        p = Path(tmpdir) / "input.xlsx"
        p.write_bytes(xlsx_bytes)
        ctx = build_context(str(p))
        missing = sheet_name_check(p)
        if missing:
            ctx.setdefault("_WARNINGS", []).append("Faltan hojas: " + ", ".join(missing))
        return ctx, missing


def generate_ppt_from_bytes(xlsx_bytes: bytes, strict: bool = False, asset_service_account_info: Optional[Dict[str, Any]] = None) -> tuple[bytes, str, list[str]]:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        input_xlsx = tmpdir_path / "input.xlsx"
        output_pptx = tmpdir_path / DEFAULT_OUTPUT_NAME
        validation_txt = tmpdir_path / "validation_report.txt"
        input_xlsx.write_bytes(xlsx_bytes)
        ctx = build_context(str(input_xlsx))
        warnings = ctx.get("_WARNINGS", [])
        if strict and warnings:
            validation_txt.write_text("VALIDACION: REVISAR\n\n" + "\n".join(f"- {w}" for w in warnings), encoding="utf-8")
            return b"", validation_txt.read_text(encoding="utf-8"), warnings
        update_ppt(str(DEFAULT_TEMPLATE), str(output_pptx), ctx, asset_service_account_info=asset_service_account_info)
        write_validation_report(validation_txt, ctx)
        return output_pptx.read_bytes(), validation_txt.read_text(encoding="utf-8"), warnings


def render_hero() -> None:
    st.markdown(
        """
        <div class="hero">
            <div class="eyebrow">Reporting automation · Maicao</div>
            <h1>Maicao Reporting Studio</h1>
            <p>Conecta la data mensual, revisa preview de KPIs, textos y piezas visuales, y exporta una presentación editable lista para compartir.</p>
            <div class="pill-row">
                <span class="pill">✨ Diseño visual ejecutivo</span>
                <span class="pill">📊 Preview antes de descargar</span>
                <span class="pill">🖼️ Top 3 + assets visuales</span>
                <span class="pill">📎 PowerPoint editable</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def kpi_card(label: str, value: str, sub: str = "") -> str:
    return f"""
    <div class="kpi-card">
        <div class="kpi-label">{escape(label)}</div>
        <div class="kpi-value">{escape(value)}</div>
        <div class="kpi-sub">{escape(sub)}</div>
    </div>
    """


def status_card(label: str, value: str, sub: str = "") -> str:
    return f"""
    <div class="status-card">
        <div class="status-label">{escape(label)}</div>
        <div class="status-value">{escape(value)}</div>
        <div class="status-sub">{escape(sub)}</div>
    </div>
    """


def text_preview_card(section: str, title: str, body: str) -> str:
    body = body if body and body != "Dato pendiente" else "Texto pendiente por completar en 15_Qualitative_Texts."
    return f"""
    <div class="text-card">
        <div class="text-section">{escape(section)}</div>
        <div class="text-title">{escape(title)}</div>
        <div class="text-body">{escape(body)}</div>
    </div>
    """


def render_status_strip(ctx: Optional[Dict[str, Any]]) -> None:
    if not ctx:
        cols = st.columns(4)
        items = [
            ("Estado", "Sin conectar", "Carga un Sheet o Excel"),
            ("Mes", "—", "Pendiente"),
            ("Validación", "—", "Pendiente"),
            ("Salida", "PPT editable", "Plantilla aprobada"),
        ]
    else:
        warnings = ctx.get("_WARNINGS", [])
        items = [
            ("Estado", "Conectado", "Data cargada correctamente"),
            ("Mes activo", ctx.get("MES", "—"), "Desde 00_Control"),
            ("Validación", "Revisar" if warnings else "OK", f"{len(warnings)} alertas" if warnings else "Sin alertas críticas"),
            ("Salida", "PPT editable", "Lista para generar"),
        ]
        cols = st.columns(4)
    for col, (label, value, sub) in zip(cols, items):
        with col:
            st.markdown(status_card(label, value, sub), unsafe_allow_html=True)


def render_overview(ctx: Dict[str, Any]) -> None:
    st.markdown('<div class="section-title">Resumen ejecutivo</div>', unsafe_allow_html=True)
    cols = st.columns(5)
    for col, (label, value_key, sub_key) in zip(cols, KPI_ORDER):
        with col:
            st.markdown(kpi_card(label, ctx.get(value_key, "—"), ctx.get(sub_key, "")), unsafe_allow_html=True)


def render_platform_table(ctx: Dict[str, Any]) -> None:
    rows = []
    for name, prefix in PLATFORM_ROWS:
        rows.append(
            {
                "Plataforma": name,
                "Views": ctx.get(f"{prefix}_VIEWS", "—"),
                "Alcance": ctx.get(f"{prefix}_REACH", "—"),
                "Interacciones": ctx.get(f"{prefix}_INTERACTIONS", "—"),
                "ER": ctx.get(f"{prefix}_ER", "—"),
                "Seguidores": ctx.get(f"{prefix}_FOLLOWERS", "—"),
            }
        )
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)


def chart_df(ctx: Dict[str, Any], key: str, labels: list[str]) -> pd.DataFrame:
    values = ctx.get("_CHART_VALUES", {}).get(key, [])
    return pd.DataFrame({"Categoría": labels[: len(values)], "Valor": values}).set_index("Categoría")


def render_chart_previews(ctx: Dict[str, Any]) -> None:
    st.markdown('<div class="section-title">Preview visual</div>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        st.caption("Visualizaciones por canal")
        st.bar_chart(chart_df(ctx, "slide4_views", ["Instagram", "Facebook", "TikTok"]), height=230)
    with c2:
        st.caption("Alcance por canal")
        st.bar_chart(chart_df(ctx, "slide4_reach", ["Instagram", "Facebook", "TikTok"]), height=230)
    with c3:
        st.caption("E.R. por canal")
        st.bar_chart(chart_df(ctx, "slide4_er", ["Instagram", "Facebook", "TikTok"]), height=230)

    c4, c5 = st.columns(2)
    with c4:
        st.caption("Squad · views por integrante")
        st.bar_chart(chart_df(ctx, "slide9_views", ["Skarleth", "Busquilla", "Cami", "Disley"]), height=260)
    with c5:
        st.caption("Squad · E.R. por integrante")
        st.bar_chart(chart_df(ctx, "slide9_er", ["Skarleth", "Busquilla", "Cami", "Disley"]), height=260)


def render_text_previews(ctx: Dict[str, Any]) -> None:
    st.markdown('<div class="section-title">Textos cualitativos que irán al PPT</div>', unsafe_allow_html=True)
    rows = []
    for section, title, key in TEXT_PREVIEW_KEYS:
        rows.append((section, title, ctx.get(key, "Dato pendiente")))
    for i in range(0, len(rows), 2):
        cols = st.columns(2)
        for col, item in zip(cols, rows[i : i + 2]):
            with col:
                st.markdown(text_preview_card(*item), unsafe_allow_html=True)




def render_media_card(item: Dict[str, Any]) -> None:
    title = item.get("content_title") or item.get("title") or "Pieza"
    url = item.get("asset_url") or item.get("image_url") or ""
    note = item.get("short_note") or item.get("note") or ""
    st.markdown('<div class="media-card">', unsafe_allow_html=True)
    if url and "..." not in str(url):
        try:
            st.image(url, use_container_width=True)
        except Exception:
            st.info("Imagen configurada; se verá en la presentación si el link es accesible.")
    else:
        st.info("Asset pendiente: agrega link de imagen en Google Sheets.")
    st.markdown(f'<div class="media-title">{escape(title)}</div>', unsafe_allow_html=True)
    metric = item.get("rank_metric_label") or ""
    value = item.get("rank_value")
    if value not in [None, ""]:
        st.markdown(f'<div class="media-meta">Ranking: {escape(metric)} · {escape(value)}</div>', unsafe_allow_html=True)
    if note:
        st.caption(str(note))
    st.markdown('</div>', unsafe_allow_html=True)


def render_media_previews(ctx: Dict[str, Any]) -> None:
    st.markdown('<div class="section-title">Contenido visual y assets</div>', unsafe_allow_html=True)
    top3 = ctx.get("_TOP3_MEDIA", {}) or {}
    if not top3:
        st.info("No se detectaron piezas para Top 3. Revisa 01_Content_Raw y 00_Control.")
        return
    for platform in ["Instagram", "Facebook", "TikTok"]:
        st.markdown(f"**Top 3 {platform}**")
        cols = st.columns(3)
        for col, item in zip(cols, top3.get(platform, []) + [{}] * 3):
            with col:
                render_media_card(item)
        st.write("")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**MMPP · imágenes del mes**")
        assets = ctx.get("_MMPP_ASSETS", []) or []
        cols = st.columns(2)
        for col, item in zip(cols, assets + [{}] * 2):
            with col:
                render_media_card(item)
    with c2:
        st.markdown("**Competencia · assets curados**")
        comp = (ctx.get("_COMPETITION_ASSETS", []) or [])[:4]
        for item in comp:
            st.caption(f"{item.get('brand','Marca')} · {item.get('display_group','Acción')}: {item.get('title','')}")

def render_validation(ctx: Optional[Dict[str, Any]]) -> None:
    st.markdown('<div class="section-title">Validación</div>', unsafe_allow_html=True)
    if not ctx:
        st.info("Conecta una fuente de datos para ver la validación del reporte.")
        return
    warnings = ctx.get("_WARNINGS", [])
    if not warnings:
        st.markdown('<span class="success-badge">✓ Reporte listo · sin alertas críticas</span>', unsafe_allow_html=True)
        st.caption("Revisa el preview y genera el PowerPoint editable cuando estés listo.")
    else:
        st.markdown(f'<span class="warn-badge">⚠ {len(warnings)} alertas por revisar</span>', unsafe_allow_html=True)
        with st.expander("Ver alertas detectadas", expanded=True):
            for w in warnings:
                st.write(f"• {w}")


def load_source_to_session(source: str, sheet_url: str, uploaded_excel: Optional[Any], sa_file: Optional[Any]) -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = Path(tmpdir)
        if source == "Google Sheets":
            sa_info = get_service_account_info(sa_file)
            st.session_state["asset_service_account_info"] = sa_info
            p = google_sheet_to_xlsx(sheet_url, sa_info, tmpdir_path / "google_sheet_input.xlsx")
            xlsx_bytes = p.read_bytes()
            source_name = "Google Sheets"
        else:
            if uploaded_excel is not None:
                p = save_uploaded_excel(uploaded_excel, tmpdir_path / "uploaded_input.xlsx")
                xlsx_bytes = p.read_bytes()
                source_name = uploaded_excel.name
            else:
                xlsx_bytes = DEFAULT_EXCEL.read_bytes()
                source_name = "Modelo incluido"
            st.session_state.setdefault("asset_service_account_info", None)
        ctx, missing_sheets = analyze_xlsx(xlsx_bytes)
        st.session_state["xlsx_bytes"] = xlsx_bytes
        st.session_state["ctx"] = ctx
        st.session_state["source_name"] = source_name
        st.session_state["missing_sheets"] = missing_sheets
        st.session_state.pop("ppt_bytes", None)
        st.session_state.pop("validation_text", None)


def render_connect_tab() -> None:
    st.markdown('<div class="section-title">Conectar fuente de datos</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="help-box">El equipo solo debe actualizar el Google Sheet maestro y conectar el link aquí. La configuración técnica queda oculta.</div>',
        unsafe_allow_html=True,
    )
    st.write("")
    source = st.radio("Fuente", ["Google Sheets", "Subir Excel"], index=0, horizontal=True, label_visibility="collapsed")

    sheet_url = ""
    uploaded_excel = None
    sa_file = None

    if source == "Google Sheets":
        sheet_url = st.text_input("Link del Google Sheet", placeholder="Pega aquí el link del Google Sheet maestro")
        if has_streamlit_secret():
            st.markdown('<span class="success-badge">✓ Conexión segura configurada</span>', unsafe_allow_html=True)
            st.caption("La app ya tiene credenciales guardadas de forma segura. No hace falta subir archivos técnicos.")
        else:
            with st.expander("Configuración avanzada", expanded=False):
                st.caption("Solo para setup local o primera configuración. No compartas esta credencial con el equipo.")
                sa_file = st.file_uploader("Credencial Google", type=["json"], help="Solo se usa durante esta sesión.")
                if sa_file is not None:
                    try:
                        info = json.loads(sa_file.getvalue().decode("utf-8"))
                        client_email = info.get("client_email")
                        if client_email:
                            st.success(f"Comparte el Google Sheet con: {client_email}")
                    except Exception:
                        st.warning("No pude leer el archivo de credencial.")
    else:
        uploaded_excel = st.file_uploader("Sube el Excel modelo", type=["xlsx"])
        with st.expander("Usar modelo incluido", expanded=False):
            st.caption("Útil para demos rápidas o pruebas locales.")
            use_default = st.checkbox("Usar archivo ejemplo incluido", value=False)
            if use_default and uploaded_excel is None:
                uploaded_excel = None

    strict = st.checkbox("Validación estricta", value=False, help="Detiene la generación si existen datos faltantes.")
    st.session_state["strict_mode"] = strict

    c1, c2 = st.columns([1, 2])
    with c1:
        if st.button("Analizar datos", type="primary", use_container_width=True):
            try:
                if source == "Google Sheets" and not sheet_url.strip():
                    raise ValueError("Pega el link del Google Sheet antes de analizar.")
                if source == "Subir Excel" and uploaded_excel is None:
                    # If no file, use the included model as a friendly default.
                    pass
                load_source_to_session(source, sheet_url, uploaded_excel, sa_file)
                st.success("Datos conectados. Revisa el preview antes de exportar.")
            except Exception as exc:
                st.error(str(exc))
    with c2:
        if st.session_state.get("ctx"):
            st.markdown(
                f'<span class="neutral-badge">Fuente activa: {escape(st.session_state.get("source_name", "—"))}</span>',
                unsafe_allow_html=True,
            )


def render_preview_tab() -> None:
    ctx = st.session_state.get("ctx")
    if not ctx:
        st.info("Primero conecta y analiza una fuente de datos en la pestaña Conectar.")
        return
    render_overview(ctx)
    st.write("")
    st.markdown('<div class="section-title">Plataformas</div>', unsafe_allow_html=True)
    render_platform_table(ctx)
    render_chart_previews(ctx)
    render_media_previews(ctx)
    render_text_previews(ctx)
    render_validation(ctx)


def render_export_tab() -> None:
    ctx = st.session_state.get("ctx")
    if not ctx:
        st.info("Conecta una fuente y revisa el preview antes de exportar.")
        return
    render_validation(ctx)
    st.write("")
    st.markdown('<div class="section-title">Crear presentación editable</div>', unsafe_allow_html=True)
    st.caption("La presentación se genera con la plantilla visual aprobada y los datos actualmente cargados.")
    if st.button("Crear PowerPoint", type="primary"):
        try:
            ppt_bytes, validation_text, warnings = generate_ppt_from_bytes(
                st.session_state["xlsx_bytes"],
                strict=bool(st.session_state.get("strict_mode", False)),
                asset_service_account_info=st.session_state.get("asset_service_account_info"),
            )
            st.session_state["ppt_bytes"] = ppt_bytes
            st.session_state["validation_text"] = validation_text
            st.session_state["export_warnings"] = warnings
            if ppt_bytes:
                st.success("Presentación creada correctamente.")
            else:
                st.warning("La presentación no se generó porque la validación estricta encontró alertas.")
        except Exception as exc:
            st.error(str(exc))

    if st.session_state.get("validation_text"):
        with st.expander("Ver reporte técnico de validación", expanded=False):
            st.text(st.session_state["validation_text"])

    if st.session_state.get("ppt_bytes"):
        st.download_button(
            "Descargar presentación editable",
            data=st.session_state["ppt_bytes"],
            file_name=DEFAULT_OUTPUT_NAME,
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            use_container_width=True,
        )


def render_help_tab() -> None:
    st.markdown('<div class="section-title">Guía rápida para el equipo</div>', unsafe_allow_html=True)
    st.markdown(
        """
        <div class="soft-card">
        <b>Flujo mensual recomendado</b><br><br>
        1. Actualizar la data en el Google Sheet maestro.<br>
        2. Completar los textos cualitativos en <code>15_Qualitative_Texts</code>.<br>
        3. Completar el plan de acción en <code>16_Action_Plan</code>.<br>
        4. Conectar el Sheet en esta app y revisar el preview.<br>
        5. Completar links de imágenes en <code>01_Content_Raw</code>, <code>21_MMPP_Assets</code> y <code>22_Competition_Assets</code>.<br>
        6. Revisar preview visual y descargar la presentación editable.<br><br>
        <span style="color:#74788A;">Tip: si algo no aparece en el preview, probablemente falta en el Google Sheet o el mes activo no coincide con <code>00_Control</code>.</span>
        </div>
        """,
        unsafe_allow_html=True,
    )


def main() -> None:
    st.set_page_config(page_title="Maicao Reporting Studio", page_icon="✨", layout="wide", initial_sidebar_state="collapsed")
    inject_css()
    render_hero()
    render_status_strip(st.session_state.get("ctx"))

    tabs = st.tabs(["Conectar", "Preview", "Exportar", "Ayuda"])
    with tabs[0]:
        render_connect_tab()
    with tabs[1]:
        render_preview_tab()
    with tabs[2]:
        render_export_tab()
    with tabs[3]:
        render_help_tab()

    st.markdown('<div class="footer-note">Maicao Reporting Studio · Google Sheets → Preview visual → PowerPoint editable</div>', unsafe_allow_html=True)


if __name__ == "__main__":
    main()

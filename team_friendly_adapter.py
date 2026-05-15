"""Team-friendly sheet adapter for Maicao Reporting Studio v14.

The team fills one sheet per slide (S01_Portada ... S12_Plan_30_Dias).
This adapter converts that simple workbook into the legacy technical model
expected by generate_report_from_template.py.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List
from openpyxl import Workbook, load_workbook

TEAM_SHEETS = [
    "S01_Portada", "S02_Dashboard", "S03_Que_Cambio", "S04_Portafolio_Canales",
    "S05_Instagram", "S06_Facebook", "S07_TikTok", "S08_TikTok_Mix", "S09_Squad",
    "S10_MMPP", "S11_Competencia", "S12_Plan_30_Dias",
]


def is_team_friendly_workbook(xlsx_path: str | Path) -> bool:
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        return "S02_Dashboard" in wb.sheetnames and "S09_Squad" in wb.sheetnames
    except Exception:
        return False


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        return v.strip()
    return v


def _fields(wb, sheet_name: str) -> Dict[str, Any]:
    if sheet_name not in wb.sheetnames:
        return {}
    ws = wb[sheet_name]
    # Find header row: Campo | Valor
    header_row = None
    for r in range(1, min(ws.max_row, 80) + 1):
        a = _clean(ws.cell(r, 1).value)
        b = _clean(ws.cell(r, 2).value)
        if str(a).lower() == "campo" and str(b).lower() == "valor":
            header_row = r
            break
    if not header_row:
        return {}
    out: Dict[str, Any] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        key = _clean(ws.cell(r, 1).value)
        if not key or str(key).startswith("TABLA:"):
            break
        out[str(key)] = _clean(ws.cell(r, 2).value)
    return out


def _table(wb, sheet_name: str, table_name: str) -> List[Dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    marker = f"TABLA: {table_name}".lower()
    start = None
    for r in range(1, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and str(val).strip().lower() == marker:
            start = r
            break
    if start is None:
        return []
    header_row = start + 1
    headers = []
    for c in range(1, ws.max_column + 1):
        h = _clean(ws.cell(header_row, c).value)
        if h in [None, ""]:
            break
        headers.append(str(h))
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        if not any(ws.cell(r, c).value not in [None, ""] for c in range(1, len(headers) + 1)):
            break
        row = {headers[c - 1]: _clean(ws.cell(r, c).value) for c in range(1, len(headers) + 1)}
        rows.append(row)
    return rows


def _write_sheet(wb, name: str, headers: list[str], rows: list[list[Any]]):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    return ws


def _f(d: Dict[str, Any], key: str, default: Any = None) -> Any:
    return d.get(key, default)


def _platform_from_sheet(team_wb, sheet_name: str, platform: str) -> Dict[str, Any]:
    f = _fields(team_wb, sheet_name)
    top = _table(team_wb, sheet_name, "Top 3")
    # Map top 3 values
    top = sorted(top, key=lambda r: float(r.get("Orden") or 999))
    rec = {
        "Mes": None,
        "Plataforma": platform,
        "Views": _f(f, "Visualizaciones"),
        "Views mes anterior": _f(f, "Visualizaciones mes anterior"),
        "Alcance": _f(f, "Alcance"),
        "Alcance mes anterior": _f(f, "Alcance mes anterior"),
        "Interacciones": _f(f, "Interacciones"),
        "Interacciones mes anterior": _f(f, "Interacciones mes anterior"),
        "ER %": _f(f, "E.R. %"),
        "ER % mes anterior": _f(f, "E.R. % mes anterior"),
        "Contenidos": _f(f, "Contenidos"),
        "Contenidos mes anterior": _f(f, "Contenidos mes anterior"),
        "Stories": _f(f, "Stories"),
        "Reels": _f(f, "Reels"),
        "Carruseles": _f(f, "Carruseles"),
        "Post": _f(f, "Post"),
        "Presupuesto": _f(f, "Presupuesto"),
        "Presupuesto mes anterior": _f(f, "Presupuesto mes anterior"),
        "Top 1 nombre": top[0].get("Nombre contenido") if len(top) > 0 else None,
        "Top 1 valor": top[0].get("Valor") if len(top) > 0 else None,
        "Top 2 nombre": top[1].get("Nombre contenido") if len(top) > 1 else None,
        "Top 2 valor": top[1].get("Valor") if len(top) > 1 else None,
        "Top 3 nombre": top[2].get("Nombre contenido") if len(top) > 2 else None,
        "Top 3 valor": top[2].get("Valor") if len(top) > 2 else None,
        "Insight corto": _f(f, "Lectura principal"),
        "Oportunidad / Nota": _f(f, "Oportunidad / optimización"),
    }
    return rec


def convert_team_friendly_to_legacy(xlsx_path: str | Path, output_path: str | Path) -> Path:
    """Convert v14 slide-by-slide input into a technical workbook and return output path."""
    xlsx_path = Path(xlsx_path)
    output_path = Path(output_path)
    if not is_team_friendly_workbook(xlsx_path):
        return xlsx_path

    src = load_workbook(xlsx_path, data_only=True)
    wb = Workbook()
    wb.remove(wb.active)

    s01 = _fields(src, "S01_Portada")
    s02 = _fields(src, "S02_Dashboard")
    s03 = _fields(src, "S03_Que_Cambio")
    s04 = _fields(src, "S04_Portafolio_Canales")
    s08 = _fields(src, "S08_TikTok_Mix")
    s09 = _fields(src, "S09_Squad")
    s10 = _fields(src, "S10_MMPP")
    s11 = _fields(src, "S11_Competencia")
    s12 = _fields(src, "S12_Plan_30_Dias")
    month = _f(s01, "Mes actual", "Marzo 2026")
    prev_month = _f(s01, "Mes anterior", "Febrero 2026")
    client = _f(s01, "Cliente", "Maicao")

    # 00_Control
    _write_sheet(wb, "00_Control", ["Campo", "Valor"], [
        ["Cliente", client],
        ["Mes actual", month],
        ["Mes anterior", prev_month],
        ["active_month", month],
        ["previous_month", prev_month],
        ["client_name", client],
        ["ig_top3_metric", "views"],
        ["fb_top3_metric", "views"],
        ["tt_top3_metric", "interactions"],
        ["enable_media_slides", "No"],
    ])

    # 09_Validaciones
    _write_sheet(wb, "09_Validaciones", ["metric", "value", "previous_value", "variation"], [
        ["views_total", _f(s02, "Visualizaciones"), _f(s02, "Visualizaciones mes anterior"), None],
        ["reach_total", _f(s02, "Alcance"), _f(s02, "Alcance mes anterior"), None],
        ["interactions_total", _f(s02, "Interacciones"), _f(s02, "Interacciones mes anterior"), None],
        ["contents_total", _f(s02, "Contenidos"), _f(s02, "Contenidos mes anterior"), None],
        ["er_total", _f(s02, "Engagement Rate"), _f(s02, "Engagement Rate mes anterior"), None],
        ["investment_total", _f(s02, "Inversión"), _f(s02, "Inversión mes anterior"), None],
    ])

    # 03_Paid_Media
    ig = _platform_from_sheet(src, "S05_Instagram", "Instagram")
    fb = _platform_from_sheet(src, "S06_Facebook", "Facebook")
    tt = _platform_from_sheet(src, "S07_TikTok", "TikTok")
    _write_sheet(wb, "03_Paid_Media", ["Mes", "Plataforma", "Monto", "Incluir Overview", "Incluir Plataforma"], [
        [month, "Overview", _f(s02, "Inversión"), "Sí", "No"],
        [month, "Instagram", ig.get("Presupuesto"), "No", "Sí"],
        [month, "Facebook", fb.get("Presupuesto"), "No", "Sí"],
        [month, "TikTok", tt.get("Presupuesto"), "No", "Sí"],
        [month, "Squad", _f(s09, "Presupuesto actual"), "No", "No"],
    ])

    # 02_Audience
    def audience_row(platform, sheet_name):
        f = _fields(src, sheet_name)
        cur = _f(f, "Seguidores")
        prev = _f(f, "Seguidores mes anterior")
        growth = None
        try:
            if cur is not None and prev is not None:
                growth = float(cur) - float(prev)
        except Exception:
            pass
        return [month, platform, cur, prev, growth, _f(f, "Mujeres %"), _f(f, "Hombres %"), _f(f, "18-24 %"), _f(f, "25-34 %"), _f(f, "35-44 %"), _f(f, "45-54 %"), _f(f, "55+ %")]
    _write_sheet(wb, "02_Audience", ["Mes", "Plataforma", "Seguidores", "Seguidores mes anterior", "Crecimiento", "Mujeres %", "Hombres %", "18-24 %", "25-34 %", "35-44 %", "45-54 %", "55+ %"], [
        audience_row("Instagram", "S05_Instagram"),
        audience_row("Facebook", "S06_Facebook"),
        audience_row("TikTok", "S07_TikTok"),
    ])

    # 13_Platform_KPIs
    headers = ["Mes","Plataforma","Views","Views mes anterior","Alcance","Alcance mes anterior","Interacciones","Interacciones mes anterior","ER %","ER % mes anterior","Contenidos","Contenidos mes anterior","Stories","Reels","Carruseles","Post","Presupuesto","Presupuesto mes anterior","Top 1 nombre","Top 1 valor","Top 2 nombre","Top 2 valor","Top 3 nombre","Top 3 valor","Insight corto","Oportunidad / Nota"]
    rows = []
    for rec in [ig, fb, tt]:
        rec["Mes"] = month
        rows.append([rec.get(h) for h in headers])
    _write_sheet(wb, "13_Platform_KPIs", headers, rows)

    # 04_Squad
    squad_rows = _table(src, "S09_Squad", "Rostros")
    _write_sheet(wb, "04_Squad", ["Mes", "Plataforma", "Talento", "Views", "Alcance", "Interacciones", "ER %", "Monto", "Reels", "Stories", "TikToks", "Comentario", "Incluir PPT"], [
        [month, "Squad", r.get("Rostro"), r.get("Views"), r.get("Alcance"), r.get("Interacciones"), r.get("E.R. %"), r.get("Presupuesto"), r.get("Reels"), r.get("Stories"), r.get("TikToks"), r.get("Lectura breve"), "Sí"] for r in squad_rows
    ])

    # 05_MMPP
    _write_sheet(wb, "05_MMPP", ["Mes", "Plataforma", "Views", "Alcance", "Interacciones", "ER %", "Contenidos", "Comentario general", "Usar en PPT"], [
        [month, "MMPP", _f(s10, "Visualizaciones"), _f(s10, "Alcance"), _f(s10, "Interacciones"), _f(s10, "E.R. %"), _f(s10, "Contenidos"), _f(s10, "Lectura general"), "Sí"]
    ])

    # 06_Competencia
    comp_rows = _table(src, "S11_Competencia", "Matriz competencia")
    _write_sheet(wb, "06_Competencia", ["Mes", "Marca", "Plataforma", "Seguidores", "Views", "Interacciones", "Lectura"], [
        [month, r.get("Marca"), r.get("Plataforma"), r.get("Seguidores"), r.get("Views"), r.get("Interacciones"), r.get("Lectura cualitativa")] for r in comp_rows
    ])

    # 01_Content_Raw from top 3 rows
    content_headers = ["month","platform","content_id","content_title","content_type","views","reach","interactions","er","asset_url","post_url","short_note","include_top3","include_platform","include_overview"]
    content_rows = []
    for platform, sheet_name in [("Instagram", "S05_Instagram"), ("Facebook", "S06_Facebook"), ("TikTok", "S07_TikTok")]:
        for r in _table(src, sheet_name, "Top 3"):
            content_rows.append([month, platform, f"{platform[:2].upper()}_{r.get('Orden')}", r.get("Nombre contenido"), "", r.get("Views"), r.get("Alcance"), r.get("Interacciones"), r.get("E.R. %"), r.get("Asset URL opcional"), "", r.get("Mini lectura"), "Sí", "Sí", "No"])
    _write_sheet(wb, "01_Content_Raw", content_headers, content_rows)

    # Qualitative texts
    qt = []
    def add_text(slide, section, placeholder, text, comment=""):
        qt.append([month, slide, section, placeholder, text, "Sí", comment])
    add_text(2, "Dashboard", "EXEC_BUSINESS_READING", _f(s02, "Lectura de negocio"))
    add_text(2, "Dashboard", "EXEC_TAG_1", _f(s02, "Tag 1"))
    add_text(2, "Dashboard", "EXEC_TAG_2", _f(s02, "Tag 2"))
    add_text(2, "Dashboard", "EXEC_TAG_3", _f(s02, "Tag 3"))
    add_text(3, "Qué cambió", "CHANGE_TITLE_1", _f(s03, "Título 1"))
    add_text(3, "Qué cambió", "CHANGE_BODY_1", _f(s03, "Texto 1"))
    add_text(3, "Qué cambió", "CHANGE_TITLE_2", _f(s03, "Título 2"))
    add_text(3, "Qué cambió", "CHANGE_BODY_2", _f(s03, "Texto 2"))
    add_text(3, "Qué cambió", "CHANGE_TITLE_3", _f(s03, "Título 3"))
    add_text(3, "Qué cambió", "CHANGE_BODY_3", _f(s03, "Texto 3"))
    add_text(3, "Qué cambió", "CHANGE_DECISION", _f(s03, "Decisión recomendada"))
    # Roles from S04 table
    channels = _table(src, "S04_Portafolio_Canales", "Canales")
    role_map = {r.get("Plataforma"): r.get("Rol") for r in channels}
    add_text(4, "Portafolio", "ROLE_INSTAGRAM", role_map.get("Instagram"))
    add_text(4, "Portafolio", "ROLE_FACEBOOK", role_map.get("Facebook"))
    add_text(4, "Portafolio", "ROLE_TIKTOK", role_map.get("TikTok"))
    add_text(4, "Portafolio", "PORTFOLIO_READING", _f(s04, "Lectura final"))
    add_text(5, "Instagram", "IG_WHAT_WORKED", _f(_fields(src, "S05_Instagram"), "Lectura principal"))
    add_text(5, "Instagram", "IG_OPTIMIZATION", _f(_fields(src, "S05_Instagram"), "Oportunidad / optimización"))
    add_text(6, "Facebook", "FB_VISUAL_READING", _f(_fields(src, "S06_Facebook"), "Lectura principal"))
    add_text(6, "Facebook", "FB_DATA_NOTE", _f(_fields(src, "S06_Facebook"), "Oportunidad / optimización"))
    add_text(7, "TikTok", "TT_CONTENT_PRINCIPLE", _f(_fields(src, "S07_TikTok"), "Lectura principal"))
    add_text(7, "TikTok", "TT_OPPORTUNITY", _f(_fields(src, "S07_TikTok"), "Oportunidad / optimización"))
    add_text(8, "TikTok Mix", "TT_MIX_READING", _f(s08, "Lectura final"))
    add_text(9, "Squad", "SQUAD_LEARNING", _f(s09, "Aprendizaje"))
    add_text(10, "MMPP", "MMPP_QUAL_READING", _f(s10, "Lectura general"))
    add_text(10, "MMPP", "MMPP_NEXT_STEP", _f(s10, "Próximo paso"))
    add_text(11, "Competencia", "COMP_OPPORTUNITY", _f(s11, "Lectura / oportunidad Maicao"))
    _write_sheet(wb, "15_Qualitative_Texts", ["Mes", "Plataforma", "Slide", "Seccion", "Placeholder", "Texto", "Usar en PPT", "Comentario"], [[r[0], "General", r[1], r[2], r[3], r[4], r[5], r[6]] for r in qt])

    # Action plan
    plan_rows = _table(src, "S12_Plan_30_Dias", "Plan de acción")
    _write_sheet(wb, "16_Action_Plan", ["Mes", "Plataforma", "Orden", "Pilar", "Accion", "Meta", "KPIs de control", "Usar en PPT"], [
        [month, "General", r.get("Orden"), r.get("Pilar"), r.get("Acción"), r.get("Meta"), _f(s12, "KPIs de control"), "Sí"] for r in plan_rows
    ])

    # Empty assets/notes with headers, so v12/v13 media code is happy.
    _write_sheet(wb, "20_Squad_Assets", ["member_name", "image_url", "order", "active"], [])
    _write_sheet(wb, "21_MMPP_Assets", ["month", "slot", "title", "image_url", "note", "active"], [])
    _write_sheet(wb, "22_Competition_Assets", ["month", "brand", "pillar", "slot", "image_url", "title", "note", "active"], [])
    _write_sheet(wb, "23_Content_Notes", ["content_id", "short_note"], [])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
